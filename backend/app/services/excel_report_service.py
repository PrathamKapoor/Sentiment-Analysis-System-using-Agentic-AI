"""Excel generation via OpenPyXL — structured worksheets with headers and
reasonable column widths, not a raw JSON dump.
"""
from datetime import datetime as dt, date, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _parse_date(value):
    """Report data arrives as ISO strings (same shape the PDF renderer
    consumes) — parse back to a real date so Excel treats the column as a
    date, not text. Never raises/fabricates: an unparseable value is left
    as-is rather than crashing report generation over a formatting nicety.
    """
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return value


def _parse_datetime(value):
    if not value:
        return None
    try:
        return dt.fromisoformat(str(value))
    except ValueError:
        return value


def _write_sheet(wb, title, headers, rows, widths=None, date_columns=None):
    """date_columns: optional {1-indexed column: number_format} applied to
    that column's data cells (header row excluded).
    """
    ws = wb.create_sheet(title=title[:31])  # Excel sheet-name length limit
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    for idx, width in enumerate(widths or [20] * len(headers), start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for col_idx, number_format in (date_columns or {}).items():
        for row_idx in range(2, len(rows) + 2):
            ws.cell(row=row_idx, column=col_idx).number_format = number_format
    return ws


def generate_excel(file_path, organisation_name, project_name, date_from, date_to, sections_data, skipped_sections):
    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    summary_rows = [
        ["Organisation", organisation_name],
        ["Project", project_name],
        ["Date range", f"{date_from} to {date_to}"],
        ["Generated", dt.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")],
        ["Report type", "AI/System Generated"],
    ]
    if skipped_sections:
        summary_rows.append(["Sections omitted (no data)", ", ".join(skipped_sections)])
    _write_sheet(wb, "Summary", ["Field", "Value"], summary_rows, widths=[28, 50])

    if "sentimentDistribution" in sections_data:
        s = sections_data["sentimentDistribution"]
        _write_sheet(wb, "Sentiment", ["Metric", "Value"], [
            ["Total reviews", s["totalReviews"]],
            ["Analysed reviews", s["analysedReviews"]],
            ["Positive count", s["positive"]["count"]],
            ["Positive %", s["positive"]["percentage"]],
            ["Negative count", s["negative"]["count"]],
            ["Negative %", s["negative"]["percentage"]],
            ["Neutral count", s["neutral"]["count"]],
            ["Neutral %", s["neutral"]["percentage"]],
            ["Winning-label score (avg, not prob.)", s["averageConfidence"]],
        ])

    if "topicAnalysis" in sections_data:
        rows = [[t["topicName"], t["reviewCount"]] for t in sections_data["topicAnalysis"]]
        _write_sheet(wb, "Topics", ["Topic", "Review Count"], rows)

    if "keywordAnalysis" in sections_data:
        rows = [
            [k["keyword"], k["frequency"], k["sentiment"], k["positiveCount"], k["negativeCount"], k["neutralCount"]]
            for k in sections_data["keywordAnalysis"]
        ]
        _write_sheet(wb, "Keywords", ["Keyword", "Frequency", "Sentiment", "Positive", "Negative", "Neutral"], rows)

    if "aspectSentiment" in sections_data:
        rows = [
            [a["name"], a["frequency"], a["positivePercentage"], a["negativePercentage"], a["neutralPercentage"], a["averageConfidence"]]
            for a in sections_data["aspectSentiment"]
        ]
        _write_sheet(wb, "Aspects", ["Aspect", "Frequency", "Positive %", "Negative %", "Neutral %", "Avg Confidence"], rows)

    if "recommendations" in sections_data:
        rows = [
            [r["text"], r["priority"], r["status"], r["aspectName"] or "", r["supportingReviewCount"]]
            for r in sections_data["recommendations"]
        ]
        _write_sheet(wb, "Recommendations", ["Recommendation", "Priority", "Status", "Aspect", "Supporting Reviews"], rows, widths=[60, 12, 12, 20, 18])

    if "alerts" in sections_data:
        rows = [
            [a["name"], a["severity"], a["status"], a["metric"], a["thresholdValue"], _parse_datetime(a["triggeredAt"])]
            for a in sections_data["alerts"]
        ]
        _write_sheet(
            wb, "Alerts", ["Alert", "Severity", "Status", "Metric", "Threshold", "Triggered At"], rows,
            date_columns={6: "YYYY-MM-DD HH:MM"},
        )

    if "representativeReviews" in sections_data:
        rows = [
            [r["text"], r["source"] or "", r["rating"], _parse_date(r["reviewDate"])]
            for r in sections_data["representativeReviews"]
        ]
        _write_sheet(
            wb, "Reviews", ["Text", "Source", "Rating", "Date"], rows, widths=[70, 20, 10, 15],
            date_columns={4: "YYYY-MM-DD"},
        )

    if "aiInterpretation" in sections_data:
        meta = sections_data["aiInterpretation"]
        source = meta.get("source", "deterministic")
        # Same labelling contract as the PDF renderer: "AI-Generated" is
        # only used when a live LLM produced the text. The deterministic
        # fallback is named as a summary, not as AI-generated.
        source_label = "AI-Generated Contextual Interpretation (LLM)" if source == "llm" \
            else "Deterministic Contextual Summary"
        rows = [
            ["Source", source_label],
            ["Provider", meta.get("provider") or ""],
            ["Model", meta.get("model") or ""],
            ["Status", meta.get("status") or ""],
            ["Latency (ms)", int(meta.get("latencyMs") or 0)],
            ["Warning", (meta.get("warning") or "")[:500]],
            ["Notice", (
                "This section was produced by a language model and is not a measured finding."
                if meta.get("source") == "llm"
                else "This section was produced deterministically from the verified metrics. No LLM was used."
            )],
            ["Interpretation", meta.get("text", "")],
        ]
        # Sheet name reflects the source so an auditor opening the file
        # sees the provenance immediately.
        sheet_title = "AI Interpretation" if meta.get("source") == "llm" else "Deterministic Summary"
        _write_sheet(wb, sheet_title, ["Field", "Value"], rows, widths=[28, 100])

    wb.save(file_path)

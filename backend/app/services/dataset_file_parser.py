import csv
import json

from app.errors.exceptions import ValidationError


def read_columns_and_rows(file_path, file_type):
    """Returns (columns: list[str], rows: list[dict]).

    ponytail: loads the whole file into memory — fine at Phase 2's scale
    (manual dataset uploads, no streaming requirement in the approved docs).
    Revisit if/when very large datasets become a real case.
    """
    if file_type == "csv":
        return _read_csv(file_path)
    if file_type == "json":
        return _read_json(file_path)
    if file_type == "excel":
        return _read_excel(file_path)
    raise ValidationError(f"Unsupported file type: {file_type}")


def _read_csv(file_path):
    with open(file_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        columns = reader.fieldnames or []
    return columns, rows


def _read_json(file_path):
    with open(file_path, encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, list):
        raise ValidationError("JSON dataset must be an array of objects")
    columns = []
    for row in data:
        if not isinstance(row, dict):
            raise ValidationError("Each JSON dataset row must be an object")
        for key in row.keys():
            if key not in columns:
                columns.append(key)
    return columns, data


def _read_excel(file_path):
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return [], []
    columns = [str(c) if c is not None else "" for c in header]
    rows = []
    for raw_row in rows_iter:
        row = {}
        for col, value in zip(columns, raw_row):
            row[col] = value
        rows.append(row)
    return columns, rows

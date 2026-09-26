import io
from datetime import date

from openpyxl import load_workbook

from tests.conftest_project import owner_context, create_project, create_reviews_directly

TODAY = date(2026, 6, 15)


def _seed_and_analyse(client, project_id, headers):
    create_reviews_directly(project_id, [
        "Delivery was terrible and late.", "Great product quality overall.",
    ], review_date=TODAY)
    client.post(f"/api/v1/projects/{project_id}/analysis/sentiment", json={}, headers=headers)


def test_topic_section_in_report_succeeds(client):
    """Regression: the topicAnalysis section used to crash report generation.

    ``topic_service.list_topics`` returns serialised dicts (``Topic.to_dict``,
    keys ``topicName`` / ``reviewCount``), but ``report_data_service`` read
    them with attribute access (``t.name``), so any report that included the
    Topics section after topic analysis had been run failed with
    ``'dict' object has no attribute 'name'``. Every existing report test
    either skipped topic analysis or omitted the section, so nothing caught it.
    """
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    create_reviews_directly(project_id, [
        "Delivery was terrible and very late.",
        "Great product quality overall, very happy.",
        "Customer service was slow to respond.",
        "Terrible experience, would not buy again.",
    ], review_date=TODAY)
    client.post(f"/api/v1/projects/{project_id}/analysis/sentiment", json={}, headers=headers)
    topics = client.post(
        f"/api/v1/projects/{project_id}/analysis/topics", json={}, headers=headers
    )
    assert topics.status_code in (200, 202)

    for file_format in ("pdf", "excel"):
        resp = client.post(
            f"/api/v1/projects/{project_id}/reports",
            json={
                "reportName": f"Topic report {file_format}",
                "dateFrom": "2026-06-01",
                "dateTo": "2026-06-30",
                "sections": ["projectOverview", "topicAnalysis"],
                "fileFormat": file_format,
            },
            headers=headers,
        )
        assert resp.status_code == 201, resp.get_json()
        body = resp.get_json()["data"]
        assert body["generationStatus"] == "complete"
        # The section must have real content, not be silently skipped.
        assert "topicAnalysis" not in (body.get("sectionsSkipped") or [])


def test_conclusion_section_in_report_succeeds(client):
    """Regression: the conclusion section used to crash report generation.

    ``sentiment_service.get_summary()`` returns each bucket as
    ``{"count", "percentage"}`` with no ``"label"`` key, but the conclusion
    section read ``dominant["label"]`` and died with ``KeyError: 'label'``
    whenever sentiment *and* aspect analysis had both been run - which is the
    default state of any analysed project. No test generated a report with
    both sections populated, so it shipped broken.
    """
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    create_reviews_directly(project_id, [
        "Delivery was terrible and very late.",
        "Great product quality overall, very happy.",
        "Customer service was slow to respond.",
        "Terrible experience, would not buy again.",
    ], review_date=TODAY)
    client.post(f"/api/v1/projects/{project_id}/analysis/sentiment", json={}, headers=headers)
    aspects = client.post(
        f"/api/v1/projects/{project_id}/analysis/aspects", json={}, headers=headers
    )
    assert aspects.status_code in (200, 202)

    for file_format in ("pdf", "excel"):
        resp = client.post(
            f"/api/v1/projects/{project_id}/reports",
            json={
                "reportName": f"Conclusion report {file_format}",
                "dateFrom": "2026-06-01",
                "dateTo": "2026-06-30",
                "sections": ["projectOverview", "conclusion"],
                "fileFormat": file_format,
            },
            headers=headers,
        )
        assert resp.status_code == 201, resp.get_json()
        body = resp.get_json()["data"]
        assert body["generationStatus"] == "complete"
        assert "conclusion" not in (body.get("sectionsSkipped") or [])


def test_conclusion_names_the_actual_dominant_sentiment_label(client, app):
    """The conclusion must name a real label, and the right one.

    Four clearly negative reviews must produce a conclusion that says
    "negative", not "positive" and not a crash.
    """
    from app.extensions import db
    from app.models import Project
    from app.services import report_data_service

    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    create_reviews_directly(project_id, [
        "Delivery was terrible and very late, I am upset.",
        "Terrible experience, would not buy again.",
        "Awful quality, broke immediately, very disappointed.",
        "Customer service was rude and unhelpful, awful experience.",
    ], review_date=TODAY)
    client.post(f"/api/v1/projects/{project_id}/analysis/sentiment", json={}, headers=headers)
    client.post(f"/api/v1/projects/{project_id}/analysis/aspects", json={}, headers=headers)

    with app.app_context():
        project = db.session.get(Project, project_id)
        # The conclusion block derives itself from these two sections, so they
        # must be gathered too.
        data, skipped = report_data_service.gather_report_data(
            project, ["sentimentDistribution", "aspectSentiment", "conclusion"], None, None
        )
        assert "conclusion" not in skipped
        conclusion = data["conclusion"]
        assert "dominant sentiment is negative" in conclusion
        assert "reviews" in conclusion


def test_topic_section_rows_use_topic_name_and_review_count(client, app):
    """The rendered section shape must match what both renderers index."""
    from app.services import report_data_service
    from app.extensions import db

    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    create_reviews_directly(project_id, [
        "Delivery was terrible and very late.",
        "Great product quality overall, very happy.",
        "Customer service was slow to respond.",
    ], review_date=TODAY)
    topics = client.post(
        f"/api/v1/projects/{project_id}/analysis/topics", json={}, headers=headers
    )
    assert topics.status_code in (200, 202)

    from app.models import Project

    with app.app_context():
        project = db.session.get(Project, project_id)
        data, skipped = report_data_service.gather_report_data(
            project, ["topicAnalysis"], None, None
        )
        assert "topicAnalysis" not in skipped
        rows = data["topicAnalysis"]
        assert rows, "expected topic rows after analysis"
        for row in rows:
            assert set(row) == {"topicName", "reviewCount"}
            assert isinstance(row["topicName"], str) and row["topicName"]
            assert isinstance(row["reviewCount"], int)


def test_generate_pdf_report(client, app):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    _seed_and_analyse(client, project_id, headers)

    resp = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={
            "reportName": "June Report", "dateFrom": "2026-06-01", "dateTo": "2026-06-30",
            "sections": ["projectOverview", "sentimentDistribution", "conclusion"], "fileFormat": "pdf",
        },
        headers=headers,
    )
    assert resp.status_code == 201
    body = resp.get_json()["data"]
    assert body["generationStatus"] == "complete"
    assert body["fileFormat"] == "pdf"

    from app.models import Report
    from app.extensions import db
    with app.app_context():
        report = db.session.get(Report, body["id"])
        import os
        assert os.path.isfile(report.file_path)
        assert os.path.getsize(report.file_path) > 0


def test_generate_excel_report_with_correct_worksheets(client, app):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    _seed_and_analyse(client, project_id, headers)
    client.post(f"/api/v1/projects/{project_id}/analysis/aspects", json={}, headers=headers)

    resp = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={
            "dateFrom": "2026-06-01", "dateTo": "2026-06-30",
            "sections": ["sentimentDistribution", "aspectSentiment"], "fileFormat": "excel",
        },
        headers=headers,
    )
    assert resp.status_code == 201
    body = resp.get_json()["data"]
    assert body["generationStatus"] == "complete"

    from app.models import Report
    from app.extensions import db
    with app.app_context():
        report = db.session.get(Report, body["id"])
        wb = load_workbook(report.file_path)
        assert "Summary" in wb.sheetnames
        assert "Sentiment" in wb.sheetnames
        assert "Aspects" in wb.sheetnames

        sentiment_sheet = wb["Sentiment"]
        header_row = [c.value for c in sentiment_sheet[1]]
        assert header_row == ["Metric", "Value"]


def test_date_range_and_sections_respected(client, app):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    _seed_and_analyse(client, project_id, headers)

    resp = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={
            "dateFrom": "2026-01-01", "dateTo": "2026-01-31",  # no data in this range
            "sections": ["sentimentDistribution"], "fileFormat": "pdf",
        },
        headers=headers,
    )
    body = resp.get_json()["data"]
    assert body["generationStatus"] == "complete"


def test_invalid_format_rejected(client):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    resp = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={"dateFrom": "2026-06-01", "dateTo": "2026-06-30", "sections": ["projectOverview"], "fileFormat": "docx"},
        headers=headers,
    )
    assert resp.status_code == 400


def test_missing_project_rejected(client):
    org_id, headers = owner_context(client)
    import uuid
    resp = client.post(
        f"/api/v1/projects/{uuid.uuid4()}/reports",
        json={"dateFrom": "2026-06-01", "dateTo": "2026-06-30", "sections": ["projectOverview"], "fileFormat": "pdf"},
        headers=headers,
    )
    assert resp.status_code == 404


def test_no_sections_rejected(client):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    resp = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={"dateFrom": "2026-06-01", "dateTo": "2026-06-30", "sections": [], "fileFormat": "pdf"},
        headers=headers,
    )
    assert resp.status_code == 400


def test_download_report(client):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    _seed_and_analyse(client, project_id, headers)
    gen = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={"dateFrom": "2026-06-01", "dateTo": "2026-06-30", "sections": ["sentimentDistribution"], "fileFormat": "pdf"},
        headers=headers,
    )
    report_id = gen.get_json()["data"]["id"]

    resp = client.get(f"/api/v1/reports/{report_id}/download", headers=headers)
    assert resp.status_code == 200
    assert resp.data[:4] == b"%PDF"


def test_cross_organisation_download_blocked(client):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    _seed_and_analyse(client, project_id, headers)
    gen = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={"dateFrom": "2026-06-01", "dateTo": "2026-06-30", "sections": ["sentimentDistribution"], "fileFormat": "pdf"},
        headers=headers,
    )
    report_id = gen.get_json()["data"]["id"]

    _, other_headers = owner_context(client, org_name="Other Org", email="other@other.test")
    resp = client.get(f"/api/v1/reports/{report_id}/download", headers=other_headers)
    assert resp.status_code == 404


def test_path_traversal_attempt_blocked(client):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)

    # No user-controlled path exists anywhere in the download route — the
    # only "path-shaped" input is the report UUID itself, which a traversal
    # string can't satisfy (it isn't a valid UUID, so access is denied before
    # any filesystem lookup happens).
    resp = client.get("/api/v1/reports/..%2F..%2F..%2Fetc%2Fpasswd/download", headers=headers)
    assert resp.status_code in (400, 404)


def test_delete_report(client):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    _seed_and_analyse(client, project_id, headers)
    gen = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={"dateFrom": "2026-06-01", "dateTo": "2026-06-30", "sections": ["sentimentDistribution"], "fileFormat": "pdf"},
        headers=headers,
    )
    report_id = gen.get_json()["data"]["id"]

    del_resp = client.delete(f"/api/v1/reports/{report_id}", headers=headers)
    assert del_resp.status_code == 200

    list_resp = client.get(f"/api/v1/projects/{project_id}/reports", headers=headers)
    assert list_resp.get_json()["data"]["items"] == []


def test_generation_failure_leaves_no_orphan_file(client, app, monkeypatch):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    _seed_and_analyse(client, project_id, headers)

    written_path = {}

    def _boom(file_path, *a, **kw):
        with open(file_path, "wb") as f:
            f.write(b"partial")  # simulate a PDF that started writing before the crash
        written_path["path"] = file_path
        raise RuntimeError("simulated renderer crash")

    import app.services.report_service as report_service_module
    monkeypatch.setattr(report_service_module, "generate_pdf", _boom)

    resp = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={"dateFrom": "2026-06-01", "dateTo": "2026-06-30", "sections": ["sentimentDistribution"], "fileFormat": "pdf"},
        headers=headers,
    )
    assert resp.status_code == 500

    from app.models import Report
    import os
    with app.app_context():
        report = Report.query.filter_by(project_id=project_id).first()
        assert report.generation_status == "failed"
        assert not report.file_path
        assert not os.path.isfile(written_path["path"])


def test_audit_logs_created(client):
    org_id, headers = owner_context(client)
    project_id = create_project(client, headers)
    _seed_and_analyse(client, project_id, headers)
    gen = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={"dateFrom": "2026-06-01", "dateTo": "2026-06-30", "sections": ["sentimentDistribution"], "fileFormat": "pdf"},
        headers=headers,
    )
    report_id = gen.get_json()["data"]["id"]
    client.get(f"/api/v1/reports/{report_id}/download", headers=headers)

    from app.models import AuditLog
    actions = {a.action for a in AuditLog.query.filter_by(entity_type="report").all()}
    assert "report.generation_started" in actions
    assert "report.generation_completed" in actions
    assert "report.downloaded" in actions
